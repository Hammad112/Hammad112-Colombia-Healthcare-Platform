"""Generate the five clinic spreadsheets the onboarding tests run against.

    python -m tests.fixtures.onboarding.generate

Every value here is invented. No file in this directory contains real patient
data, and none may: see CLAUDE.md rule 7.

The fixtures escalate deliberately, from a file any importer would manage to
one built to break it:

    1_clean_ips.xlsx        a tidy export. The control: it must map with no model call.
    2_receptionist.xlsx     a hand-made workbook: title rows, a two-row header,
                            merged cells, hidden rows and columns, a trailing total,
                            a second table on the same sheet, a stale hidden sheet.
    3_excel_csv_es.csv      what Spanish Excel writes: CP1252, semicolons, comma
                            decimals, an embedded newline, doubled quotes, a ragged row.
    4_corrupted.xlsx        what Excel does to clinical data left to its own devices:
                            cédulas as numbers, a document number in scientific
                            notation, an undecidable date column, serial dates,
                            SI/NO booleans, times as floats and as "8:30 a. m.".
    5_hostile.xlsx          a file that attacks the parser.

Sentinel values (`SENTINEL_*`) are planted in every patient-data cell of
fixtures 1-4. The payload-safety test asserts that not one of them ever appears
in anything sent to a model provider, which is what proves the mapping stage
transmits no patient data.
"""

from __future__ import annotations

import csv
import datetime as dt
import shutil
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).parent

# Planted in patient-data cells so a test can prove they never leave the machine.
SENTINEL_NAME = "Zzyzx Sentinelensen Marcadorez"
SENTINEL_DOCUMENT = "9999888877"
SENTINEL_PHONE = "3009998877"
SENTINEL_EMAIL = "sentinel.marcador@example.invalid"
SENTINELS = (SENTINEL_NAME, SENTINEL_DOCUMENT, SENTINEL_PHONE, SENTINEL_EMAIL)

# Invented people. Colombian shape: two given names, two surnames.
_PATIENTS = [
    ("CC", "1045678901", "Carlos Andrés", "Pérez Gómez", "3001234567", "Sura"),
    ("CC", "1023456789", "María Fernanda", "López Torres", "3119876543", "Nueva EPS"),
    ("CC", "71234567", "Luis Ernesto", "Castro Vargas", "3204567890", "EPS Sanitas"),
    ("CC", "43567890", "Diana Carolina", "Muñoz Restrepo", "3156789012", "Coosalud"),
    ("TI", "1102345678", "Luisa Fernanda", "Ortiz Mejía", "3067891234", "Nueva EPS"),
    ("CC", "1067891234", "Andrés Felipe", "Suárez Vélez", "3098765432", "Salud Total"),
    # A defunct EPS: historical rows name insurers that no longer exist, and the
    # importer must accept them rather than reject a real patient.
    ("CC", "1054321987", "Sandra Milena", "Zuluaga Franco", "3145678901", "Medimás"),
    ("CE", "E0456789", "Jean Pierre", "Dubois Martín", "3213456789", "PARTICULAR"),
    ("CC", "1076543210", "Juan Pablo", "Henao Correa", "3178901234", "Compensar"),
    (SENTINEL_NAME, SENTINEL_DOCUMENT, SENTINEL_PHONE, SENTINEL_EMAIL, "", ""),
]

_DOCTORS = [
    ("M01", "Dra. Patricia Gómez", "Ortopedia", "12"),
    ("M02", "Dr. Juan Ramírez", "Ortopedia", "14"),
    ("M03", "Dr. Carlos Vega", "Cirugía General", "08"),
    ("M04", "Dra. Lina Marín", "Ginecología", "15"),
]


def _autosize(ws: Worksheet) -> None:
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(width + 2, 40)


# --------------------------------------------------------------------- fixture 1
def clean_ips() -> Path:
    """A tidy export. Headers are ordinary Spanish; values are already clean.

    This is the control: the deterministic matcher must map every column with no
    model call at all.
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Pacientes"
    ws.append(
        [
            "Tipo Documento",
            "Número Documento",
            "Nombres",
            "Apellidos",
            "Fecha Nacimiento",
            "Celular",
            "Correo Electrónico",
            "EPS",
        ]
    )
    for index, (doc_type, doc, given, family, phone, eps) in enumerate(_PATIENTS[:-1]):
        ws.append(
            [
                doc_type,
                doc,
                given,
                family,
                # A real date cell, not text: the easy case.
                dt.date(1980 + index, (index % 12) + 1, (index % 27) + 1),
                phone,
                f"{given.split()[0].lower()}.{family.split()[0].lower()}@example.invalid",
                eps,
            ]
        )
    # The sentinel patient, so fixture 1 also participates in the privacy test.
    ws.append(
        [
            "CC",
            SENTINEL_DOCUMENT,
            SENTINEL_NAME.split()[0],
            " ".join(SENTINEL_NAME.split()[1:]),
            dt.date(1991, 4, 3),
            SENTINEL_PHONE,
            SENTINEL_EMAIL,
            "Sura",
        ]
    )
    _autosize(ws)

    doctors = wb.create_sheet("Medicos")
    doctors.append(["Código", "Profesional", "Especialidad", "Consultorio"])
    for row in _DOCTORS:
        doctors.append(list(row))
    _autosize(doctors)

    appointments = wb.create_sheet("Citas")
    appointments.append(
        ["Documento Paciente", "Profesional", "Fecha Cita", "Hora", "Estado", "Tipo de Consulta"]
    )
    for index, patient in enumerate(_PATIENTS[:6]):
        appointments.append(
            [
                patient[1],
                _DOCTORS[index % len(_DOCTORS)][0],
                dt.date(2026, 10, (index % 28) + 1),
                dt.time(7 + index, 15 * (index % 4)),
                ["Confirmada", "Pendiente", "Atendida", "Cancelada"][index % 4],
                "Primera vez" if index % 2 else "Control",
            ]
        )
    _autosize(appointments)

    path = HERE / "1_clean_ips.xlsx"
    wb.save(path)
    return path


# --------------------------------------------------------------------- fixture 2
def receptionist() -> Path:
    """A workbook a receptionist maintains by hand.

    Most Colombian providers have fewer than ten professionals, so this shape is
    more representative than a system export: a title above the data, a header
    split across two rows, merged cells, hidden rows and columns, a totals row,
    and a second table pasted below the first.
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "AGENDA"

    ws["A1"] = "CLÍNICA SAN RAFAEL — AGENDA DE CITAS"
    ws.merge_cells("A1:H1")
    ws["A2"] = "Generado el 30/09/2026"
    ws.merge_cells("A2:H2")
    ws["A3"] = None

    # A two-row header: "DATOS DEL PACIENTE" spans the first three columns and
    # the real names live underneath it.
    ws["A4"] = "DATOS DEL PACIENTE"
    ws.merge_cells("A4:C4")
    ws["D4"] = "CITA"
    ws.merge_cells("D4:F4")
    ws["G4"] = "CONTACTO"
    ws.merge_cells("G4:H4")
    ws.append([])  # keeps openpyxl's row pointer honest
    for column, value in enumerate(
        [
            "TIPO DOC",
            "IDENTIFICACION",
            "NOMBRE COMPLETO",
            "FECHA",
            "HORA",
            "MEDICO",
            "CELULAR",
            "EPS",
        ],
        start=1,
    ):
        ws.cell(row=5, column=column, value=value)

    row = 6
    for index, (doc_type, doc, given, family, phone, eps) in enumerate(_PATIENTS[:-1]):
        ws.cell(row=row, column=1, value={"CC": "C.C.", "TI": "T.I.", "CE": "C.E."}[doc_type])
        ws.cell(row=row, column=2, value=doc)
        # One column for the whole name, the common hand-made shape.
        ws.cell(row=row, column=3, value=f"{given} {family}")
        ws.cell(row=row, column=4, value=f"{(index % 28) + 1:02d}/10/2026")
        ws.cell(row=row, column=5, value=f"{7 + index}:{15 * (index % 4):02d}")
        ws.cell(row=row, column=6, value=_DOCTORS[index % len(_DOCTORS)][1])
        ws.cell(row=row, column=7, value=phone)
        ws.cell(row=row, column=8, value=eps)
        row += 1

    # Two hidden rows holding real patients. They must be imported, and flagged:
    # hiding a row is not a deletion, and silently dropping them loses patients.
    for _doc_type, doc, given, family, phone, eps in (
        ("CC", "1032165498", "Claudia Patricia", "Escobar Rendón", "3178901299", "Sura"),
        ("CC", SENTINEL_DOCUMENT, *SENTINEL_NAME.split(maxsplit=1), SENTINEL_PHONE, "Nueva EPS"),
    ):
        ws.cell(row=row, column=1, value="C.C.")
        ws.cell(row=row, column=2, value=doc)
        ws.cell(row=row, column=3, value=f"{given} {family}")
        ws.cell(row=row, column=4, value="15/10/2026")
        ws.cell(row=row, column=5, value="09:30")
        ws.cell(row=row, column=6, value=_DOCTORS[0][1])
        ws.cell(row=row, column=7, value=phone)
        ws.cell(row=row, column=8, value=eps)
        ws.row_dimensions[row].hidden = True
        row += 1

    # A trailing junk row: not a patient, and it must not become one.
    ws.cell(row=row + 1, column=1, value="TOTAL:")
    ws.cell(row=row + 1, column=2, value=len(_PATIENTS) + 1)

    # A second table pasted below the first, with its own header.
    second = row + 4
    ws.cell(row=second, column=1, value="PACIENTES EN LISTA DE ESPERA")
    ws.cell(row=second + 1, column=1, value="IDENTIFICACION")
    ws.cell(row=second + 1, column=2, value="NOMBRE")
    ws.cell(row=second + 1, column=3, value="TELEFONO")
    ws.cell(row=second + 2, column=1, value="1088776655")
    ws.cell(row=second + 2, column=2, value="Rosa Elena Quintero Díaz")
    ws.cell(row=second + 2, column=3, value="3112223344")

    # A hidden column holding the code the clinic really keys on.
    ws.cell(row=5, column=10, value="COD_EPS_INTERNO")
    for offset in range(len(_PATIENTS) - 1):
        ws.cell(row=6 + offset, column=10, value=f"EPS{offset:03d}")
    ws.column_dimensions["J"].hidden = True

    # A stale hidden sheet from a previous month.
    old = wb.create_sheet("AGOSTO (viejo)")
    old.append(["IDENTIFICACION", "NOMBRE", "FECHA"])
    old.append(["1045678901", "Carlos Andrés Pérez Gómez", "12/08/2026"])
    old.sheet_state = "hidden"

    path = HERE / "2_receptionist.xlsx"
    wb.save(path)
    return path


# --------------------------------------------------------------------- fixture 3
def excel_csv_es() -> Path:
    """What Excel writes on a Spanish-locale Windows machine.

    Semicolon-delimited, CP1252-encoded, comma decimal separators. A sniffer that
    assumes commas reads this as one column per line, which is why the reader
    scores delimiters instead of guessing.
    """
    rows = [
        [
            "TIPO DOC",
            "IDENTIFICACION",
            "NOMBRE COMPLETO",
            "CELULAR",
            "EPS",
            "VALOR COPAGO",
            "DIRECCION",
        ],
        [
            "C.C.",
            "1045678901",
            "Carlos Andrés Pérez Gómez",
            "3001234567",
            "Sura",
            "1.250,50",
            "Calle 10 # 5-20",
        ],
        [
            "C.C.",
            "1023456789",
            'María Fernanda "Mafe" López',
            "3119876543",
            "Nueva EPS",
            "0,00",
            "Cra 7 # 12-34",
        ],
        # An embedded newline inside a quoted field.
        [
            "T.I.",
            "1102345678",
            "Luisa Fernanda Ortiz Mejía",
            "3067891234",
            "Nueva EPS",
            "15.300,75",
            "Av. Siempre Viva 742\nApto 301",
        ],
        [
            "C.C.",
            SENTINEL_DOCUMENT,
            SENTINEL_NAME,
            SENTINEL_PHONE,
            "Coosalud",
            "2.000,00",
            "Calle 1",
        ],
    ]
    path = HERE / "3_excel_csv_es.csv"
    with path.open("w", encoding="cp1252", newline="") as handle:
        writer = csv.writer(handle, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        writer.writerows(rows)
        # A ragged row: fewer fields than the header promises.
        handle.write("C.C.;71234567;Luis Ernesto Castro Vargas\r\n")

    # The same file with a BOM, to prove utf-8-sig is tried before utf-8.
    bom = HERE / "3_excel_csv_es_bom.csv"
    with bom.open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle, delimiter=";").writerows(rows)
    return path


# --------------------------------------------------------------------- fixture 4
def corrupted() -> Path:
    """What Excel does to clinical data when nobody stops it.

    Each column here is a decision the importer has to get right, and every one
    of them is either "reject, the value is gone" or "ask, the value is
    ambiguous". None of them may be silently repaired.
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Datos"
    ws.append(
        [
            "IDENTIFICACION",  # numeric: leading zeros already destroyed
            "DOCUMENTO_LARGO",  # scientific notation: digits unrecoverable
            "NOMBRE COMPLETO",
            "FECHA_NACIMIENTO",  # every value <= 12: DD/MM vs MM/DD undecidable
            "FECHA_CITA",  # has a value > 12: decidable as day-first
            "FECHA_SERIAL",  # raw Excel serial numbers
            "HORA",  # "8:30 a. m." with a non-breaking space
            "HORA_FRACCION",  # time as a fraction of a day
            "ASISTIO",  # SI / Sí / X / 1 / blank / N/A
            "TELEFONO",
        ]
    )

    rows: list[list[object]] = [
        # Leading zero gone: the cell is a number, so 0123456 became 123456.
        [
            123456,
            1.23457e11,
            "Carlos Andrés Pérez Gómez",
            "03/04/1991",
            "15/10/2026",
            46300,
            "8:30 a. m.",
            0.354166666,
            "SI",
            "3001234567",
        ],
        [
            1023456789,
            1.02345e11,
            "María Fernanda López Torres",
            "05/06/1985",
            "22/10/2026",
            45900,
            "2:30 p. m.",
            0.604166666,
            "Sí",
            "+57 311 987 6543",
        ],
        [
            71234567,
            9.87654e10,
            "Luis Ernesto Castro Vargas",
            "11/12/1990",
            "03/11/2026",
            46000,
            "10:00 a. m.",
            0.416666666,
            "X",
            "3204567890",
        ],
        [
            43567890,
            4.35679e10,
            "Diana Carolina Muñoz Restrepo",
            "07/08/1979",
            "28/10/2026",
            46100,
            "11:15 a. m.",
            0.46875,
            1,
            "315 678 9012",
        ],
        [
            1102345678,
            1.10235e11,
            SENTINEL_NAME,
            "02/03/1995",
            "09/11/2026",
            46200,
            "9:45 a. m.",
            0.40625,
            None,
            SENTINEL_PHONE,
        ],
        # 8 digits starting with 3: a mobile that lost a digit. Unrecoverable.
        [
            1067891234,
            1.06789e11,
            "Andrés Felipe Suárez Vélez",
            "10/11/1988",
            "17/11/2026",
            46400,
            "3:00 p. m.",
            0.625,
            "N/A",
            "30987654",
        ],
    ]
    for row in rows:
        ws.append(row)

    # Non-breaking (U+00A0) and narrow no-break (U+202F) spaces, as Excel writes
    # the Spanish 12-hour form. A plain str.split() on " " misses both.
    ws["G2"] = "8:30 a. m."
    ws["G3"] = "2:30 p. m."

    _autosize(ws)
    path = HERE / "4_corrupted.xlsx"
    wb.save(path)
    return path


# --------------------------------------------------------------------- fixture 5
def hostile() -> tuple[Path, ...]:
    """Files built to break the parser rather than to be imported.

    Nothing here should ever reach the mapping stage: the reader must refuse
    each one, and the subprocess limits must contain anything that slips past.
    """
    produced: list[Path] = []

    # A formula that spreadsheet software may execute when the error report is
    # reopened. Stored as an ordinary string in a patient-name column.
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Pacientes"
    ws.append(["IDENTIFICACION", "NOMBRE", "TELEFONO"])
    ws.append(["1045678901", "=cmd|'/c calc'!A1", "3001234567"])
    ws.append(["1023456789", "+cmd|'/c calc'!A1", "3119876543"])
    ws.append(["71234567", "-2+3+cmd|'/c calc'!A1", "3204567890"])
    ws.append(["43567890", "@SUM(1+1)*cmd|'/c calc'!A1", "3156789012"])
    ws.append(["1102345678", "\tDATA()", "3067891234"])
    injection = HERE / "5_formula_injection.xlsx"
    wb.save(injection)
    produced.append(injection)

    # A declared dimension of one cell over a sheet that really has 5,000 rows.
    # openpyxl's read-only mode trusts the declaration and returns one cell,
    # which silently truncates the import unless reset_dimensions() is called.
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Pacientes"
    ws.append(["IDENTIFICACION", "NOMBRE"])
    for index in range(5000):
        ws.append([f"10{index:08d}", f"Paciente Número {index}"])
    lying = HERE / "5_lying_dimension.xlsx"
    wb.save(lying)
    _rewrite_dimension(lying, "A1:A1")
    produced.append(lying)

    # A zip bomb: a small file that expands to roughly a gigabyte.
    bomb = HERE / "5_zip_bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
        archive.writestr("xl/sharedStrings.xml", "A" * (1024 * 1024 * 1024))
    produced.append(bomb)

    # An XXE attempt: an external entity pointed at a local file.
    xxe = HERE / "5_xxe.xlsx"
    with zipfile.ZipFile(xxe, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
        archive.writestr("xl/sharedStrings.xml", _XXE_SHARED_STRINGS)
    produced.append(xxe)

    # An executable renamed to .csv: the magic bytes disagree with the extension.
    disguised = HERE / "5_not_really_csv.csv"
    disguised.write_bytes(b"MZ\x90\x00\x03" + b"\x00" * 200)
    produced.append(disguised)

    return tuple(produced)


_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>'
)

_XXE_SHARED_STRINGS = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<!DOCTYPE sst [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
    '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    "<si><t>&xxe;</t></si></sst>"
)


def _rewrite_dimension(path: Path, dimension: str) -> None:
    """Replace the declared sheet dimension, leaving the real rows in place."""
    temporary = path.with_suffix(".tmp")
    with (
        zipfile.ZipFile(path) as source,
        zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                start = text.find("<dimension")
                if start != -1:
                    end = text.find("/>", start) + 2
                    text = f'{text[:start]}<dimension ref="{dimension}"/>{text[end:]}'
                data = text.encode("utf-8")
            target.writestr(item, data)
    shutil.move(temporary, path)


def main() -> None:
    produced = [clean_ips(), receptionist(), excel_csv_es(), corrupted(), *hostile()]
    for path in produced:
        print(f"  {path.name:32} {path.stat().st_size:>10,} bytes")
    print(f"{len(produced)} fixture files written to {HERE}")


if __name__ == "__main__":
    main()
